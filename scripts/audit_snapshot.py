"""Read-only release checks for a completed full-pull snapshot."""
import argparse
import json
import re
import sys
import importlib
import pkgutil
from collections import Counter
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from common.config import OUTPUT_COLUMNS, TEMPLATE_SHEET
from common.excel_utils import _find_header_row, _build_template_header_maps, _resolve_template_col
from common.validation import validate_event_data, parse_event_date
import scrapers.play365


def provider_classes():
    labels = {}
    for module in pkgutil.iter_modules(scrapers.play365.__path__):
        imported = importlib.import_module(f'scrapers.play365.{module.name}')
        for value in vars(imported).values():
            if isinstance(value, type) and value.__module__ == imported.__name__:
                labels[f'{value.SCRAPER_NAME} (Play365Base)'] = value.__name__
    return labels


def destination(root, run_date, location):
    match = re.match(r'^(.*?)\s*\((.*?)\)\s*$', location.strip())
    region, organization = match.groups() if match else (location, 'DefaultOrg')
    clean = lambda value: re.sub(r'[^\w-]', '_', value.strip())
    return root / run_date / clean(organization) / f'{clean(region)}_{run_date}.xlsx'


def audit(manifest_path, ledger_path):
    manifest = json.loads(Path(manifest_path).read_text(encoding='utf-8'))
    ledger = json.loads(Path(ledger_path).read_text(encoding='utf-8'))['events']
    problems, summaries, empty_locations = [], [], []
    exported_identities = set()
    output_root = Path(manifest['output_root']).resolve()
    today = date.fromisoformat(manifest['run_date'])
    class_names = provider_classes()
    if manifest['mode'] != 'snapshot':
        problems.append('Release audit requires snapshot mode')
    if manifest['status'] != 'complete':
        problems.append(f"Manifest status is {manifest['status']}")
    for provider, item in manifest['providers'].items():
        if item['status'] != 'complete':
            problems.append(f"Provider {provider}: {item['status']}: {item.get('error', '')}")
        for location in item['locations']:
            if provider == 'Play365BaseScraper' and location in manifest['providers']:
                continue
            if f'{provider}|{location}' not in manifest['locations']:
                problems.append(f'Missing location result: {provider}|{location}')
    for key, item in manifest['locations'].items():
        provider, location = key.split('|', 1)
        if item['status'] != 'complete':
            problems.append(f"Location {key}: {item['status']}: {item.get('collection_error', '')}")
        for metric in ('validation_failures', 'extraction_failures', 'persistence_failures'):
            if item.get(metric, 0):
                problems.append(f"{key}: {metric}={item[metric]}")
        path = destination(output_root, manifest['run_date'], location)
        if not path.exists():
            if item.get('collection_complete') and item['rows_written'] == 0 and item.get('candidates_processed', 0) == 0:
                empty_locations.append(key)
                continue
            problems.append(f'Missing snapshot workbook: {path}')
            continue
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = workbook[TEMPLATE_SHEET]
            header_row, headers, _ = _find_header_row(sheet, OUTPUT_COLUMNS, search_rows=min(10, sheet.max_row))
            maps = _build_template_header_maps(headers)
            columns = {name: _resolve_template_col(name, *maps) for name in OUTPUT_COLUMNS}
            if not all(columns.values()):
                problems.append(f'{key}: incomplete 56-column schema')
                continue
            rows = {}
            for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                if not any(value not in (None, '') for value in cells):
                    continue
                row = {name: cells[column - 1] if column <= len(cells) else None for name, column in columns.items()}
                rows[row_number] = row
                valid, error = validate_event_data(dict(row))
                if not valid:
                    problems.append(f'{key}: import row {row_number}: {error}')
                start = parse_event_date(row['Start-date'])
                if start is None or start <= today:
                    problems.append(f'{key}: import row {row_number} is not strictly future')
                for bound, comparison in (('start_date', lambda a, b: a < b), ('end_date', lambda a, b: a > b)):
                    if manifest.get(bound) and start and comparison(start, date.fromisoformat(manifest[bound])):
                        problems.append(f'{key}: import row {row_number} exceeds {bound}')
            identities, indexed_rows = [], []
            if '_ScraperExportIndex' in workbook.sheetnames:
                index = workbook['_ScraperExportIndex']
                if index.sheet_state != 'hidden':
                    problems.append(f'{key}: export index is visible')
                for identity, row_number, url, name in index.iter_rows(min_row=2, max_col=4, values_only=True):
                    if identity is None:
                        continue
                    identities.append(identity)
                    exported_identities.add(identity)
                    indexed_rows.append(row_number)
                    row = rows.get(row_number)
                    if row is None or row.get('Website-url') != url or row.get('Tournament-name') != name:
                        problems.append(f'{key}: export index does not match import row {row_number}')
                    entry = ledger.get(identity)
                    if entry is None:
                        entry = next((e for e in ledger.values() if e.get('export_identity') == identity), None)
                    if entry is None or entry.get('scraper') != class_names.get(provider, provider) or location not in entry.get('locations', [entry.get('location')]):
                        problems.append(f'{key}: missing ledger identity/location membership {identity}')
            elif rows:
                problems.append(f'{key}: populated workbook has no identity index')
            if len(set(identities)) != len(identities) or len(set(indexed_rows)) != len(indexed_rows):
                problems.append(f'{key}: duplicate export identity or import row')
            if set(indexed_rows) != set(rows):
                problems.append(f'{key}: unindexed import rows')
            if len(rows) != item['rows_written']:
                problems.append(f"{key}: workbook rows {len(rows)} != manifest rows {item['rows_written']}")
            summaries.append({'location': key, 'rows': len(rows), 'file': str(path)})
        except Exception as error:
            problems.append(f'{key}: workbook audit failed: {error}')
        finally:
            workbook.close()
    for identity, entry in ledger.items():
        file = entry.get('export_file')
        if file and Path(file).resolve().is_relative_to(output_root) and (entry.get('export_status') != 'exported' or entry.get('export_error')):
            problems.append(f'Unacknowledged export: {identity}')
    return {'run_id': manifest['run_id'], 'passed': not problems, 'workbooks': len(summaries),
            'source_incomplete_skipped': sum(item.get('source_incomplete_skipped', 0) for item in manifest['locations'].values()),
            'verified_empty_locations_without_workbooks': len(empty_locations),
            'rows': sum(item['rows'] for item in summaries),
            'unique_events': len(exported_identities),
            'provider_statuses': dict(Counter(p['status'] for p in manifest['providers'].values())),
            'problems': problems, 'locations': summaries}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('manifest')
    parser.add_argument('--ledger', default='events.json')
    parser.add_argument('--report', required=True)
    args = parser.parse_args()
    report = audit(args.manifest, args.ledger)
    Path(args.report).write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps({key: value for key, value in report.items() if key not in ('locations', 'problems')}))
    print(f"Problems: {len(report['problems'])}; report: {args.report}")
    sys.exit(0 if report['passed'] else 1)
