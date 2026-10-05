"""Compare all 56 indexed import fields in same-date snapshot copies, read-only.

Usage: python scripts/compare_snapshot_fields.py BEFORE AFTER REPORT.json
Exit 1 flags removed records or changed fields for source-backed review.
Additions are reported separately. Only compare copies after their writers stop.
"""
import sys
import json
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from common.config import OUTPUT_COLUMNS, TEMPLATE_SHEET
from common.excel_utils import _find_header_row, _build_template_header_maps, _resolve_template_col


def rows(root):
    result = {}
    root = Path(root)
    for path in root.rglob('*.xlsx'):
        with_workbook = load_workbook(path, read_only=True)
        try:
            ws = with_workbook[TEMPLATE_SHEET]
            header, headings, _ = _find_header_row(ws, OUTPUT_COLUMNS, search_rows=min(10, ws.max_row))
            maps = _build_template_header_maps(headings)
            columns = [_resolve_template_col(c, *maps) for c in OUTPUT_COLUMNS]
            if not all(columns):
                raise ValueError('Incomplete schema: ' + str(path))
            data = {n: [cells[c-1] if c <= len(cells) and cells[c-1] is not None else '' for c in columns]
                    for n, cells in enumerate(ws.iter_rows(min_row=header+1, values_only=True), header+1)}
            if '_ScraperExportIndex' not in with_workbook.sheetnames:
                if any(any(v != '' for v in row) for row in data.values()):
                    raise ValueError('Missing identity index: ' + str(path))
                continue
            for key, number, *_ in with_workbook['_ScraperExportIndex'].iter_rows(min_row=2, values_only=True):
                if key is not None:
                    identity = (str(path.relative_to(root)), key)
                    if identity in result:
                        raise ValueError('Duplicate identity')
                    result[identity] = data[number]
        finally:
            with_workbook.close()
    return result


if __name__ == '__main__':
    before, after = rows(sys.argv[1]), rows(sys.argv[2])
    removed = sorted(set(before) - set(after))
    added = sorted(set(after) - set(before))
    changed = [{'file': k[0], 'identity': k[1], 'fields':
                {c: {'before': a, 'after': b} for c, a, b in zip(OUTPUT_COLUMNS, before[k], after[k]) if a != b}}
               for k in sorted(set(before) & set(after)) if before[k] != after[k]]
    report = {'before_rows': len(before), 'after_rows': len(after), 'unchanged_rows': len(set(before) & set(after))-len(changed),
              'removed': removed, 'added': added, 'changed': changed}
    Path(sys.argv[3]).write_text(json.dumps(report, indent=2), encoding='utf-8')
    print({**{k:v for k,v in report.items() if k not in ['removed','added','changed']},
           'removed': len(removed), 'added': len(added), 'changed': len(changed)})
    sys.exit(bool(removed or changed))
