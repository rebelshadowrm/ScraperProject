import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from common.config import OUTPUT_COLUMNS, TEMPLATE_SHEET
from common.identity import event_identity, normalize_url, discovered_identities
from common.json_utils import default_event_entry, load_events
from common.pull_options import PullOptions
from common.runner import process_location
from common.validation import validate_event_data
from scrapers.topgun_scraper import TopGunScraper
from test_pull_recovery import FakeScraper, payload


class IdentityTests(unittest.TestCase):
    def test_known_unlisted_future_event_is_rechecked_only_for_snapshot(self):
        from scrapers.pg_scraper import PGScraper
        self.assertTrue(PGScraper.REFRESH_KNOWN_FUTURE_EVENTS)
        for mode in ('snapshot', 'incremental'):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as temporary:
                root, scraper = Path(temporary), FakeScraper()
                scraper.REFRESH_KNOWN_FUTURE_EVENTS = True
                url = 'https://example.com/known'
                key = event_identity(scraper, url)
                entry = default_event_entry(url=url, location='Texas')
                entry.update(scraper='FakeScraper', processed=True, export_status='exported', event_data={**payload(url, 'Cached title'), 'Phone-number': '111-111-1111'})
                events = {'events': {key: entry}}
                for suffix, changes in (('past', {'event_data': {**payload(url), 'Start-date': '01/01/2000', 'End-date': '01/02/2000'}}), ('deleted', {'is_deleted': True})):
                    other = dict(entry, url=url + '/' + suffix, **changes)
                    events['events'][event_identity(scraper, other['url'])] = other
                fresh = {**payload(url, 'Fresh published title'), 'Phone-number': '222-222-2222'}
                with patch.object(scraper, 'collect_event_links', return_value=[]), patch.object(scraper, 'extract_event_details', return_value=fresh) as extract:
                    process_location(location_name='Texas', base_url='https://example.com', scraper=scraper, driver=None,
                                     events_data=events, failed_data={'failed_events': []}, output_columns=OUTPUT_COLUMNS,
                                     json_file=str(root / 'events.json'), failed_events_file=str(root / 'failed.json'), download_dir='',
                                     shutdown_check=lambda: False, output_base_folder=str(root / 'Sheets'),
                                     scheduler_history_file=str(root / 'history.json'), pull_options=PullOptions(mode))
                self.assertEqual(extract.call_count, 1 if mode == 'snapshot' else 0)
                actual = load_events(root / 'events.json')['events'][key]['event_data']
                self.assertEqual(actual['Tournament-name'], 'Fresh published title' if mode == 'snapshot' else 'Cached title')
                self.assertEqual(actual['Phone-number'], '222-222-2222' if mode == 'snapshot' else '111-111-1111')

    def test_unlisted_future_source_failure_cannot_certify_cached_data(self):
        with tempfile.TemporaryDirectory() as temporary:
            root, scraper = Path(temporary), FakeScraper()
            scraper.REFRESH_KNOWN_FUTURE_EVENTS = True
            url = 'https://example.com/known'
            entry = default_event_entry(url=url, location='Texas')
            entry.update(scraper='FakeScraper', processed=True, export_status='exported', event_data=payload(url, 'Cached title'))
            with patch.object(scraper, 'collect_event_links', return_value=[]), patch.object(scraper, 'extract_event_details', return_value=None):
                process_location(location_name='Texas', base_url='https://example.com', scraper=scraper, driver=None,
                                 events_data={'events': {event_identity(scraper, url): entry}}, failed_data={'failed_events': []},
                                 output_columns=OUTPUT_COLUMNS, json_file=str(root / 'events.json'), failed_events_file=str(root / 'failed.json'),
                                 download_dir='', shutdown_check=lambda: False, output_base_folder=str(root / 'Sheets'),
                                 scheduler_history_file=str(root / 'history.json'), pull_options=PullOptions('snapshot'))
            self.assertEqual(scraper._last_location_result['status'], 'partial')
            self.assertEqual(scraper._last_location_result['extraction_failures'], 1)

    def test_duplicate_listing_links_do_not_reset_legacy_processed_state(self):
        scraper = FakeScraper()
        url = "https://example.com/Case"
        entry = default_event_entry(url=url, location="Texas")
        entry["processed"] = True
        events = {"events": {url: entry}}
        identities = discovered_identities(events, {"failed_events": []}, scraper, [url, url], "Texas")
        key = event_identity(scraper, url)
        self.assertEqual(identities, {key})
        self.assertTrue(events["events"][key]["processed"])
        self.assertNotIn(url, events["events"])

    def topgun(self):
        scraper = TopGunScraper()
        scraper._filter_loaded = True
        scraper._state_to_location_ids = {"TX": [1], "OK": [2]}
        return scraper

    def tournament(self, identity, name):
        return dict(TournamentID=identity, TournamentName=name, SportID=1, LocationID=1,
                    TournamentTypeID=2, StartDate="2099-01-01T00:00:00Z", EndDate="2099-01-02T00:00:00Z")

    def test_same_date_location_and_type_retain_both_tournaments(self):
        scraper = self.topgun()
        with patch.object(scraper, "_get", return_value=[self.tournament(1, "A"), self.tournament(2, "B")]):
            links = scraper.collect_event_links(None, "https://topgunstats.com/tournaments?state=TX")
        self.assertEqual(len(set(links)), 2)
        self.assertEqual(len({event_identity(scraper, u) for u in links}), 2)
        self.assertEqual({t["TournamentName"] for t in scraper.tournaments_by_url.values()}, {"A", "B"})

    def test_tournament_identity_survives_date_change(self):
        scraper = self.topgun()
        with patch.object(scraper, "_get", return_value=[self.tournament(1, "A")]):
            first = scraper.collect_event_links(None, "https://topgunstats.com/tournaments?state=TX")[0]
        scraper._tournaments_cache[0]["StartDate"] = "2099-02-01T00:00:00Z"
        second = scraper.collect_event_links(None, "https://topgunstats.com/tournaments?state=TX")[0]
        self.assertNotEqual(first, second)
        self.assertEqual(event_identity(scraper, first), event_identity(scraper, second))

    def test_umpire_type_name_is_not_overwritten(self):
        scraper = TopGunScraper()
        with patch.object(scraper, "_get", return_value={"locations": [], "types": [
                {"TournamentTypeID": 2, "TournamentTypeName": "Umpire"}]}):
            scraper._load_filter_data()
        self.assertEqual(scraper._type_id_to_name[2], "Umpire")
        self.assertIn(2, scraper._banned_type_ids)

    def test_url_normalization_preserves_path_and_query_case(self):
        self.assertEqual(normalize_url("HTTPS://Example.COM/Case?token=AbC"),
                         "https://example.com/Case?token=AbC")
        self.assertNotEqual(normalize_url("https://example.com/Case"), normalize_url("https://example.com/case"))

    def test_reversed_dates_are_rejected(self):
        event = payload("https://example.com")
        event["End-date"] = "12/31/2098"
        self.assertFalse(validate_event_data(event)[0])

    def test_snapshot_refreshes_processed_event_and_replaces_stale_rows(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            scraper = FakeScraper()
            url = "https://example.com/Case"
            events = {"events": {url: default_event_entry(url=url, location="Texas")}}
            events["events"][url].update(processed=True, event_data=payload(url, "Stale"))
            kwargs = dict(location_name="Texas", base_url="https://example.com", scraper=scraper, driver=None,
                          events_data=events, failed_data={"failed_events": []}, output_columns=OUTPUT_COLUMNS,
                          json_file=str(root / "events.json"), failed_events_file=str(root / "failed.json"),
                          download_dir="", shutdown_check=lambda: False, output_base_folder=str(root / "Sheets"),
                          scheduler_history_file=str(root / "history.json"), pull_options=PullOptions("snapshot"),
                          run_date="2080-01-01")
            with patch.object(scraper, "collect_event_links", return_value=[url]):
                process_location(**kwargs)
            self.assertEqual(scraper.extractions, [url])
            entry = load_events(root / "events.json")["events"][event_identity(scraper, url)]
            self.assertEqual(entry["event_data"]["Tournament-name"], "Tournament")
            self.assertEqual(entry["export_identity"], url)
            workbook_path = next((root / "Sheets").rglob("*.xlsx"))
            with patch.object(scraper, "collect_event_links", return_value=[]):
                process_location(**kwargs)
            workbook = load_workbook(workbook_path)
            try:
                rows = [r for r in workbook[TEMPLATE_SHEET].iter_rows(min_row=3, values_only=True) if r[0]]
                self.assertEqual(rows, [])
            finally:
                workbook.close()

    def test_date_scope_excludes_events_outside_selected_season(self):
        scope = PullOptions("snapshot", date(2098, 1, 1), date(2098, 12, 31))
        self.assertFalse(scope.includes(payload("https://example.com")))
        with self.assertRaises(ValueError):
            PullOptions("snapshot", date(2099, 1, 1), date(2098, 1, 1))


if __name__ == "__main__":
    unittest.main()
