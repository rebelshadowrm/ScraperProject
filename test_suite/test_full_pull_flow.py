import logging
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from common.config import TEMPLATE_SHEET
from common.json_utils import load_events, save_events, default_event_entry
from common.run_manifest import RunManifest
from test_pull_recovery import FakeScraper, payload
from ui import cli


class FullPullFlowTests(unittest.TestCase):
    def test_staged_export_is_recovered_when_dynamic_location_discovery_fails(self):
        class GhostScraper(FakeScraper):
            SCRAPER_NAME = "Ghost"
            def uses_dynamic_locations(self):
                return True
            def get_dynamic_locations(self):
                raise RuntimeError("dynamic listing unavailable")
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            url = "https://example.com/event"
            location = "Texas (Ghost)"
            output = root / "recovered.xlsx"
            entry = default_event_entry(url=url, location=location, scraper="GhostScraper",
                                        discovered_from="https://example.com/source")
            entry.update(export_status="pending", event_data=payload(url), export_file=str(output), export_location=location)
            save_events({"events": {"GhostScraper|" + url: entry}}, root / "events.json")
            with patch.multiple(cli, JSON_FILE=str(root / "events.json"), FAILED_EVENTS_FILE=str(root / "failed.json"),
                                SCHEDULER_HISTORY_FILE=str(root / "history.json"), shutdown_flag=False), \
                    patch.object(cli, "RunManifest", side_effect=lambda options, **kwargs: RunManifest(options, root=root)), \
                    patch.object(cli, "get_scraper_classes", return_value=[GhostScraper]), \
                    patch.object(cli, "load_locations", return_value={}), \
                    patch.object(cli, "setup_logger", return_value=logging.getLogger("fixture")), \
                    patch.object(cli, "build_run_log_path", return_value=str(root / "fixture.log")), \
                    patch.object(cli, "flush_logs"), \
                    patch.object(GhostScraper, "extract_event_details", side_effect=AssertionError("unexpected extraction")) as extract:
                result = cli._run_scrape_flow(all_scrapers=True)
            self.assertEqual(result.data["status"], "partial")
            self.assertIn("dynamic listing unavailable", result.data["providers"]["GhostScraper"]["error"])
            self.assertEqual(result.data["providers"]["GhostScraper"]["locations"], [location])
            self.assertTrue(next(iter(load_events(root / "events.json")["events"].values()))["processed"])
            self.assertTrue(output.exists())
            extract.assert_not_called()

    def test_real_cli_pipeline_refreshes_fixture_and_reports_partial_coverage(self):
        class FixtureScraper(FakeScraper):
            SCRAPER_NAME = "Fixture"
            def collect_event_links(self, driver, base_url, **kwargs):
                if base_url.endswith("missing"):
                    raise RuntimeError("fixture listing unavailable")
                return ["https://example.com/fixture"]

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            def manifest(options, **kwargs):
                return RunManifest(options, root=root)
            with patch.multiple(cli, JSON_FILE=str(root / "events.json"), FAILED_EVENTS_FILE=str(root / "failed.json"),
                                SCHEDULER_HISTORY_FILE=str(root / "history.json"), shutdown_flag=False), \
                    patch.object(cli, "RunManifest", side_effect=manifest), \
                    patch.object(cli, "get_scraper_classes", return_value=[FixtureScraper]), \
                    patch.object(cli, "load_locations", return_value={"Texas (Fixture)": "https://example.com/listing",
                                                                         "Oklahoma (Fixture)": "https://example.com/missing"}), \
                    patch.object(cli, "setup_logger", return_value=logging.getLogger("fixture")), \
                    patch.object(cli, "build_run_log_path", return_value=str(root / "fixture.log")), \
                    patch.object(cli, "flush_logs"), patch("builtins.input", side_effect=AssertionError("unexpected prompt")):
                result = cli._run_scrape_flow(all_scrapers=True)
                second = cli._run_scrape_flow(all_scrapers=True)
            self.assertEqual(result.data["mode"], "snapshot")
            self.assertEqual(result.data["status"], "partial")
            self.assertEqual(result.data["locations"]["FixtureScraper|Oklahoma (Fixture)"]["collection_complete"], False)
            self.assertNotEqual(result.output_root, second.output_root)
            entry = next(iter(load_events(root / "events.json")["events"].values()))
            self.assertEqual(entry["attempt_count"], 2)
            self.assertTrue(entry["processed"])
            self.assertTrue(Path(result.path).exists())
            self.assertTrue(Path(second.path).exists())
            outputs = list(Path(second.output_root).rglob("*.xlsx"))
            self.assertEqual(len(outputs), 1)
            workbook = load_workbook(outputs[0])
            try:
                rows = [r for r in workbook[TEMPLATE_SHEET].iter_rows(min_row=3, values_only=True) if r[0]]
                self.assertEqual(len(rows), 1)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
