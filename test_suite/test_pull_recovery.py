import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from common import batch_runner
from common.config import OUTPUT_COLUMNS, TEMPLATE_SHEET
from common.excel_utils import append_rows_to_template, create_workbook_from_template
from common.json_utils import default_event_entry, load_events, save_events
from common.runner import process_location
from common.scheduler_history import load_scheduler_history
from common.storage import ledger_lock, save_json_atomic
from common.pull_options import PullOptions


def payload(url, name="Tournament"):
    return {"Tournament-name": name, "Start-date": "01/01/2099", "End-date": "01/02/2099",
            "State-FullName": "Texas", "Website-url": url}


class FakeScraper:
    REQUIRES_DRIVER = False
    BATCH_DELAY_SECONDS = 0

    def __init__(self):
        self.extractions = []

    def extract_event_details(self, url, **kwargs):
        self.extractions.append(url)
        return payload(url)

    def collect_event_links(self, *args, **kwargs):
        return []


class RecoveryTests(unittest.TestCase):
    def test_recovered_driver_attempt_is_reported_without_unresolved_failure(self):
        from selenium.common.exceptions import InvalidSessionIdException
        self.scraper.REQUIRES_DRIVER = True
        self.urls = self.urls[:1]
        with patch.object(self.scraper, 'extract_event_details', side_effect=[InvalidSessionIdException('inactive session'), payload(self.urls[0])]), \
                patch.object(batch_runner, 'restart_driver', return_value='recovered-driver') as restart:
            driver, written, metrics = self.run_batch(driver='old-driver')
        self.assertEqual(driver, 'recovered-driver')
        self.assertEqual(written, 1)
        self.assertEqual(metrics['events_attempted'], 2)
        self.assertEqual(metrics['driver_restarts'], 1)
        self.assertEqual(metrics['extraction_failures'], 0)
        self.assertEqual(self.failed['failed_events'], [])
        restart.assert_called_once()

    def test_failed_driver_restart_remains_an_unresolved_failure(self):
        from selenium.common.exceptions import InvalidSessionIdException
        self.scraper.REQUIRES_DRIVER = True
        self.urls = self.urls[:1]
        with patch.object(self.scraper, 'extract_event_details', side_effect=InvalidSessionIdException('inactive session')), \
                patch.object(batch_runner, 'restart_driver', side_effect=RuntimeError('cannot restart')):
            _, written, metrics = self.run_batch(driver='old-driver')
        self.assertEqual(written, 0)
        self.assertEqual(metrics['extraction_failures'], 1)
        self.assertEqual(metrics['driver_restarts'], 0)
        self.assertIn('Driver restart failed', self.failed['failed_events'][0]['reason'])

    def test_explicit_source_exclusion_is_durable_and_retryable(self):
        from common.source_exclusions import SourceEventIncomplete
        with patch.object(self.scraper, 'extract_event_details', side_effect=SourceEventIncomplete('DATE TBD')):
            _, written, metrics = self.run_batch()
        self.assertEqual(written, 0)
        self.assertEqual(metrics['source_incomplete_skipped'], 2)
        self.assertEqual(metrics['extraction_failures'], 0)
        self.assertEqual(metrics['validation_failures'], 0)
        for entry in load_events(self.ledger)['events'].values():
            self.assertFalse(entry['processed'])
            self.assertEqual(entry['export_status'], 'source_incomplete')
            self.assertEqual(entry['source_exclusion'], 'DATE TBD')
        _, written, metrics = self.run_batch()
        self.assertEqual(written, 2)
        self.assertEqual(metrics['source_incomplete_skipped'], 0)
        for entry in load_events(self.ledger)['events'].values():
            self.assertTrue(entry['processed'])
            self.assertEqual(entry['export_status'], 'exported')
            self.assertNotIn('source_exclusion', entry)

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.ledger = str(self.root / "events.json")
        self.failed_path = str(self.root / "failed.json")
        self.output = str(self.root / "sheet.xlsx")
        self.urls = [f"https://example.com/event/{n}" for n in range(2)]
        self.events = {"events": {u: default_event_entry(url=u, location="Texas") for u in self.urls}}
        self.failed = {"failed_events": []}
        self.scraper = FakeScraper()

    def run_batch(self, **overrides):
        arguments = dict(scraper=self.scraper, driver=None,
                         urls_to_process=[{"name": u, "url": u} for u in self.urls],
                         events_data=self.events, failed_data=self.failed, json_file=self.ledger,
                         failed_events_file=self.failed_path, output_file=self.output,
                         columns=OUTPUT_COLUMNS, selected_location="Texas", download_dir="",
                         shutdown_check=lambda: False, batch_size=1)
        arguments.update(overrides)
        return batch_runner.process_events_in_batches(**arguments)

    def rows(self, path=None):
        workbook = load_workbook(path or self.output)
        try:
            return [r for r in workbook[TEMPLATE_SHEET].iter_rows(min_row=3, values_only=True) if r[0]]
        finally:
            workbook.close()

    def test_checkpoint_failure_after_export_replays_without_duplicates(self):
        calls = 0

        def checkpoint(data, path):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise OSError("checkpoint unavailable after export")
            save_events(data, path)

        with patch.object(batch_runner, "save_events", side_effect=checkpoint):
            _, written, metrics = self.run_batch()
        self.assertEqual(written, 1)
        self.assertEqual(metrics["persistence_failures"], 1)
        self.assertFalse(self.events["events"][self.urls[0]]["processed"])
        self.events = load_events(self.ledger)
        self.run_batch()
        self.assertEqual(len(self.rows()), 2)
        self.assertEqual(self.scraper.extractions, self.urls)
        self.assertTrue(all(e["processed"] for e in load_events(self.ledger)["events"].values()))

    def test_failed_workbook_preserves_payload_and_reports_failure(self):
        with patch.object(batch_runner, "append_rows_to_template", side_effect=PermissionError("locked")):
            _, written, metrics = self.run_batch()
        self.assertEqual(written, 0)
        self.assertEqual(metrics["persistence_failures"], 1)
        entry = load_events(self.ledger)["events"][self.urls[0]]
        self.assertEqual(entry["export_status"], "pending")
        self.assertEqual(entry["event_data"]["Website-url"], self.urls[0])
        self.assertEqual(entry["event_data"]["Tournament-name"], "Tournament")
        self.assertIn("Workbook export failed", self.failed["failed_events"][0]["reason"])

    def test_staging_failure_never_writes_workbook(self):
        with patch.object(batch_runner, "save_events", side_effect=OSError("disk full")), \
                patch.object(batch_runner, "append_rows_to_template") as export:
            _, _, metrics = self.run_batch()
        export.assert_not_called()
        self.assertEqual(metrics["persistence_failures"], 1)

    def test_retry_keeps_original_destination_on_another_date(self):
        with patch.object(batch_runner, "append_rows_to_template", side_effect=PermissionError("locked")):
            self.run_batch(urls_to_process=[{"name": self.urls[0], "url": self.urls[0]}])
        self.events = load_events(self.ledger)
        next_day = str(self.root / "next-day.xlsx")
        self.run_batch(output_file=next_day, urls_to_process=[{"name": self.urls[0], "url": self.urls[0]}])
        self.assertEqual(len(self.rows()), 1)
        self.assertFalse(Path(next_day).exists())

    def test_pending_export_recovers_without_listing_rediscovery(self):
        with patch.object(batch_runner, "append_rows_to_template", side_effect=PermissionError("locked")):
            self.run_batch(urls_to_process=[{"name": self.urls[0], "url": self.urls[0]}])
        process_location(location_name="Texas", base_url="https://example.com", scraper=self.scraper,
                         driver=None, events_data=load_events(self.ledger), failed_data=self.failed,
                         output_columns=OUTPUT_COLUMNS, json_file=self.ledger,
                         failed_events_file=self.failed_path, download_dir="", shutdown_check=lambda: False,
                         output_base_folder=str(self.root / "Sheets"),
                         scheduler_history_file=str(self.root / "history.json"))
        self.assertEqual(len(self.rows()), 1)
        self.assertEqual(self.scraper.extractions, [self.urls[0]])

    def test_interruption_flushes_already_extracted_rows(self):
        def extract(url, **kwargs):
            if url == self.urls[1]:
                raise KeyboardInterrupt
            return payload(url)
        with patch.object(self.scraper, "extract_event_details", side_effect=extract):
            _, written, metrics = self.run_batch(batch_size=2)
        self.assertEqual(written, 1)
        self.assertTrue(metrics["interrupted"])
        self.assertEqual(len(self.rows()), 1)

    def test_repeated_export_updates_same_identity(self):
        create_workbook_from_template(self.output)
        append_rows_to_template(self.output, [payload(self.urls[0], "Before")], OUTPUT_COLUMNS, ["key"])
        append_rows_to_template(self.output, [payload(self.urls[0], "After")], OUTPUT_COLUMNS, ["key"])
        self.assertEqual(len(self.rows()), 1)
        self.assertEqual(self.rows()[0][0], "After")

    def test_existing_legacy_row_is_adopted_without_duplicate(self):
        create_workbook_from_template(self.output)
        append_rows_to_template(self.output, [payload(self.urls[0])], OUTPUT_COLUMNS)
        append_rows_to_template(self.output, [payload(self.urls[0])], OUTPUT_COLUMNS, ["key"])
        self.assertEqual(len(self.rows()), 1)

    def test_failed_json_serialization_preserves_previous_checkpoint(self):
        save_json_atomic({"good": True}, self.ledger)
        with self.assertRaises(TypeError):
            save_json_atomic({"bad": object()}, self.ledger)
        self.assertEqual(json.loads(Path(self.ledger).read_text()), {"good": True})
        self.assertFalse(list(self.root.glob(".events-*")))

    def test_atomic_workbook_replace_failure_preserves_previous_rows(self):
        create_workbook_from_template(self.output)
        with patch("common.storage.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                append_rows_to_template(self.output, [payload(self.urls[0])], OUTPUT_COLUMNS, ["key"])
        self.assertEqual(self.rows(), [])

    def test_ledger_lock_blocks_second_writer_and_releases(self):
        with ledger_lock(self.ledger):
            with self.assertRaises(RuntimeError):
                with ledger_lock(self.ledger):
                    self.fail("second writer acquired the lock")
        with ledger_lock(self.ledger):
            pass

    def test_history_reports_persistence_failure(self):
        history = str(self.root / "history.json")
        with patch.object(self.scraper, "collect_event_links", return_value=[self.urls[0]]), \
                patch.object(batch_runner, "append_rows_to_template", side_effect=PermissionError("locked")):
            process_location(location_name="Texas", base_url="https://example.com", scraper=self.scraper,
                             driver=None, events_data=self.events, failed_data=self.failed,
                             output_columns=OUTPUT_COLUMNS, json_file=self.ledger,
                             failed_events_file=self.failed_path, download_dir="", shutdown_check=lambda: False,
                             output_base_folder=str(self.root / "Sheets"), scheduler_history_file=history)
        entry = next(iter(load_scheduler_history(history)["locations"].values()))
        self.assertTrue(entry["last_failed"])
        self.assertEqual(entry["recent_runs"][-1]["persistence_failures"], 1)

    def test_snapshot_exports_once_after_multiple_durable_staging_batches(self):
        real_export = append_rows_to_template
        with patch.object(batch_runner, "append_rows_to_template", wraps=real_export) as export, \
                patch.object(batch_runner, "save_events", wraps=save_events) as checkpoint:
            _, written, metrics = self.run_batch(pull_options=PullOptions("snapshot"), batch_size=1)
        self.assertEqual(written, 2)
        self.assertEqual(export.call_count, 1)
        self.assertEqual(checkpoint.call_count, 3)  # Two staging commits, one acknowledgement.
        self.assertEqual(metrics["batches_attempted"], 2)
        self.assertEqual(len(self.rows()), 2)

    def test_snapshot_acknowledgement_failure_replays_all_staged_rows_once(self):
        calls = 0
        def checkpoint(data, path):
            nonlocal calls
            calls += 1
            if calls == 3:
                raise OSError("snapshot acknowledgement unavailable")
            save_events(data, path)
        with patch.object(batch_runner, "save_events", side_effect=checkpoint):
            _, written, metrics = self.run_batch(pull_options=PullOptions("snapshot"), batch_size=1)
        self.assertEqual(written, 2)
        self.assertEqual(metrics["persistence_failures"], 1)
        self.events = load_events(self.ledger)
        self.run_batch(pull_options=PullOptions("snapshot"), batch_size=1)
        self.assertEqual(len(self.rows()), 2)
        self.assertEqual(self.scraper.extractions, self.urls)

    def test_export_beyond_template_rows_preserves_styles_and_dropdowns(self):
        create_workbook_from_template(self.output)
        workbook = load_workbook(self.output)
        sheet = workbook[TEMPLATE_SHEET]
        bottom = sheet.max_row
        sheet.cell(bottom, 1).value = "Existing last template row"
        workbook.save(self.output)
        workbook.close()
        append_rows_to_template(self.output, [payload(self.urls[0])], OUTPUT_COLUMNS, ["key"])
        workbook = load_workbook(self.output)
        try:
            sheet = workbook[TEMPLATE_SHEET]
            self.assertEqual(sheet.cell(bottom + 1, 1).value, "Tournament")
            self.assertEqual(sheet.cell(bottom + 1, 1)._style, sheet.cell(3, 1)._style)
            self.assertTrue(any(f"K{bottom + 1}" in validation for validation in sheet.data_validations.dataValidation))
        finally:
            workbook.close()

    def test_distinct_identities_can_share_the_same_public_url(self):
        create_workbook_from_template(self.output)
        append_rows_to_template(self.output, [payload(self.urls[0], "A"), payload(self.urls[0], "B")],
                                OUTPUT_COLUMNS, ["tournament-a", "tournament-b"])
        append_rows_to_template(self.output, [payload(self.urls[0], "B changed")], OUTPUT_COLUMNS, ["tournament-b"])
        self.assertEqual([row[0] for row in self.rows()], ["A", "B changed"])

    def test_date_scope_skips_are_reported_separately_from_expired_events(self):
        _, written, metrics = self.run_batch(pull_options=PullOptions("snapshot", end_date=date(2098, 12, 31)))
        self.assertEqual(written, 0)
        self.assertEqual(metrics["date_scope_skipped"], 2)
        self.assertEqual(metrics["non_future_skipped"], 0)
        self.assertFalse(Path(self.output).exists())

    def test_snapshot_recovers_old_destination_then_refreshes_into_new_snapshot(self):
        url = self.urls[0]
        with patch.object(batch_runner, "append_rows_to_template", side_effect=PermissionError("locked")):
            self.run_batch(urls_to_process=[{"name": url, "url": url}])
        with patch.object(self.scraper, "collect_event_links", return_value=[url]), \
                patch.object(self.scraper, "extract_event_details", return_value=payload(url, "Fresh details")) as extract:
            process_location(location_name="Texas", base_url="https://example.com", scraper=self.scraper,
                             driver=None, events_data=load_events(self.ledger), failed_data=self.failed,
                             output_columns=OUTPUT_COLUMNS, json_file=self.ledger,
                             failed_events_file=self.failed_path, download_dir="", shutdown_check=lambda: False,
                             output_base_folder=str(self.root / "Snapshot"),
                             scheduler_history_file=str(self.root / "history.json"), pull_options=PullOptions("snapshot"))
        self.assertEqual(self.rows()[0][0], "Tournament")
        new_snapshot = next((self.root / "Snapshot").rglob("*.xlsx"))
        self.assertEqual(self.rows(new_snapshot)[0][0], "Fresh details")
        extract.assert_called_once()


if __name__ == "__main__":
    unittest.main()
